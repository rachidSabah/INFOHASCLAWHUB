from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def create_resume():
    doc = Document()
    
    # Page setup for A4
    section = doc.sections[0]
    section.page_height = Inches(11.69)
    section.page_width = Inches(8.27)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)

    # Style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    # Header
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("OUSSAMA EL FATIMI")
    run.bold = True
    run.font.size = Pt(18)
    
    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.add_run("Rabat, Morocco | +212 7 08 69 03 70 | oussamaelfatimi29@gmail.com\n")
    contact.add_run("Dubai, UAE (Ready to Relocate) | Willing to Wear Uniform | Shift Flexible")
    contact.paragraph_format.space_after = Pt(12)

    def add_section_header(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12)

    # Professional Summary
    add_section_header("PROFESSIONAL SUMMARY")
    summary = doc.add_paragraph(
        "Highly motivated and trilingual Customer Service Professional with specialized training in Aviation & Hospitality from INFOHAS. "
        "Demonstrated expertise in passenger handling at Rabat International Airport, including check-in, boarding operations, and conflict resolution. "
        "Adept at maintaining high safety and security standards while delivering premium service in multicultural environments. "
        "Committed to the Emirates 'Fly Better' philosophy, striving to exceed passenger expectations through professionalism, empathy, and cultural awareness."
    )

    # Core Competencies
    add_section_header("CORE COMPETENCIES")
    competencies = [
        ("Aviation Services:", "Passenger Check-in & Boarding, Baggage Services, STEB Compliance, Immigration Procedures."),
        ("Customer Excellence:", "Service Recovery, Conflict Resolution, VIP & PRM Assistance, Premium Hospitality."),
        ("Communication:", "Trilingual (English, French, Arabic), Interpersonal Diplomacy, Team Collaboration."),
        ("Technical Skills:", "Microsoft Office Suite, CRM Platforms, Data Entry, Fast Typing."),
        ("Soft Skills:", "Cultural Sensitivity, Stress Management, Problem Solving, Grooming Excellence.")
    ]
    for category, skills in competencies:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(category)
        run.bold = True
        p.add_run(f" {skills}")
        p.paragraph_format.space_after = Pt(2)

    # Professional Experience
    add_section_header("PROFESSIONAL EXPERIENCE")
    
    # Job 1
    p = doc.add_paragraph()
    run = p.add_run("Rabat International Airport")
    run.bold = True
    p.add_run(" | Customer Service Agent Intern")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    date_p = doc.add_paragraph("May 2024 – October 2024")
    date_p.paragraph_format.space_after = Pt(4)

    bullets = [
        "Facilitated seamless passenger processing for international flights, managing check-in counters and boarding gate operations for 200+ passengers daily.",
        "Ensured 100% compliance with airport safety, security, and immigration protocols to maintain operational integrity.",
        "Provided dedicated assistance to passengers with reduced mobility (PRM) and unaccompanied minors, enhancing their travel experience.",
        "Coordinated effectively with ground handling and baggage services to resolve passenger inquiries and minimize delays.",
        "Applied active listening and service recovery techniques to handle passenger concerns with professionalism and empathy."
    ]
    for bullet in bullets:
        doc.add_paragraph(bullet, style='List Bullet').paragraph_format.space_after = Pt(2)

    # Job 2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    run = p.add_run("Gold Shop — Premium Retail")
    run.bold = True
    p.add_run(" | Sales Attendant")
    
    date_p2 = doc.add_paragraph("January 2022 – November 2023")
    date_p2.paragraph_format.space_after = Pt(4)

    bullets2 = [
        "Delivered high-end customer service in a luxury retail environment, consistently exceeding sales targets through the F.A.B. (Features, Advantages, Benefits) method.",
        "Managed high-value financial transactions and inventory with strict adherence to Standard Operating Procedures (SOPs).",
        "Curated visual merchandising displays to reflect brand standards and attract international clientele.",
        "Built long-term customer relationships by providing personalized shopping experiences and expert product knowledge."
    ]
    for bullet in bullets2:
        doc.add_paragraph(bullet, style='List Bullet').paragraph_format.space_after = Pt(2)

    # Education
    add_section_header("EDUCATION")
    p = doc.add_paragraph()
    run = p.add_run("Diploma in Hospitality & Aviation")
    run.bold = True
    p.add_run(" | INFOHAS, Morocco")
    doc.add_paragraph("Focus: Aviation Safety & Security, CRM, Hospitality Management, Customer Service Excellence.", style='List Bullet')
    
    p = doc.add_paragraph()
    run = p.add_run("High School Certificate")
    run.bold = True
    p.add_run(" | Morocco")

    # Additional Information
    add_section_header("ADDITIONAL INFORMATION")
    info = [
        "Languages: English (Fluent), French (Fluent), Arabic (Native).",
        "Availability: Immediate availability for international relocation and shift-based work (24/7).",
        "Standards: Fully committed to Emirates Group grooming and uniform standards.",
        "Technical: Proficient in Microsoft Word, Excel, and Outlook."
    ]
    for item in info:
        doc.add_paragraph(item, style='List Bullet').paragraph_format.space_after = Pt(2)

    # Ensure it saves in the project directory
    project_dir = r"C:\Users\piopi\Documents\gemini-projects\gemini-desktop"
    filename = "Oussama_El_Fatimi_Resume_Emirates_Optimized.docx"
    save_path = os.path.join(project_dir, filename)
    doc.save(save_path)
    print(f"Resume generated successfully at: {save_path}")

if __name__ == "__main__":
    create_resume()
